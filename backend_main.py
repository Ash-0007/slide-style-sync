import sys
import os
import io
import shutil
import uuid
import json
from datetime import datetime
from typing import Optional, Dict, List, Any, Tuple
from pathlib import Path
import tempfile
import traceback

import aiofiles
from fastapi import FastAPI, File, UploadFile, HTTPException, Form, Depends
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse, JSONResponse, FileResponse, Response
from pydantic import BaseModel, Field
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from pptx.exc import PackageNotFoundError


# --- Image Processing Imports (Aspose - Removed as unused) ---
# try:
#     import aspose.slides as slides
#     import aspose.pydrawing as drawing
#     from PIL import Image as PILImage # Keep PIL if needed elsewhere, otherwise remove
# except ImportError:
#     # print("ERROR: Required libraries 'aspose.slides' and/or 'Pillow' not found.")
#     # print("Please install them using: pip install aspose.slides Pillow")
#     slides = None
#     drawing = None
#     PILImage = None

# --- Pillow Import (If needed independently) ---
# Make sure Pillow is still needed, e.g., by pptx_tools or direct use
# If not, remove this too and from requirements.txt
try:
    from PIL import Image as PILImage 
except ImportError:
    print("ERROR: Required library 'Pillow' not found.")
    print("Please install it using: pip install Pillow")
    PILImage = None




save_pptx_as_png = None 
try:
    
    from pptx_tools.utils import save_pptx_as_png as pptx_save_func
    save_pptx_as_png = pptx_save_func 
    print("Successfully imported pptx_tools.utils.save_pptx_as_png")
except ImportError:
    print("ERROR: Required library 'python-pptx-tools' not found.")
    print("Please install it using: pip install python-pptx-tools")
    
except Exception as pptx_tools_e:
    print(f"ERROR: Could not import from pptx_tools: {pptx_tools_e}")
    print("Ensure any dependencies (like COM components or PowerPoint) are available.")




SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
IMAGES_FOLDER = os.path.join(SCRIPT_DIR, "images_server")
TEMP_FOLDER = os.path.join(SCRIPT_DIR, "temp_files")
PROFILES_FILE = os.path.join(SCRIPT_DIR, "user_profiles.json")


os.makedirs(IMAGES_FOLDER, exist_ok=True)
os.makedirs(TEMP_FOLDER, exist_ok=True)



def load_profiles() -> Dict[str, str]:
    """Loads profile data (name -> filename) from the JSON file."""
    try:
        if os.path.exists(PROFILES_FILE):
            with open(PROFILES_FILE, 'r') as f:
                data = json.load(f)
                
                return data if isinstance(data, dict) else {}
        return {}
    except (json.JSONDecodeError, IOError) as e:
        print(f"Error loading profiles file ({PROFILES_FILE}): {e}. Returning empty profiles.")
        return {}

def save_profiles(profiles: Dict[str, str]):
    """Saves profile data (name -> filename) to the JSON file."""
    try:
        with open(PROFILES_FILE, 'w') as f:
            json.dump(profiles, f, indent=4)
    except IOError as e:
        print(f"Error saving profiles file ({PROFILES_FILE}): {e}")




DEFAULT_SEARCH_TEXTS = {
    'theme': 'Speak friend and enter',
    'day': 'WED',
    'date': '30th',
    'month': 'APRIL',
    'year': '2025',
    'tmod': 'TM TMOD NAME',
    'ge': 'GE Name',
    'speaker1': 'TM Person1',
    'speaker2': 'TM Person2'
    
}

IMAGE_SHAPE_NAMES = {
    'tmod_image': 'TMOD_Image_Shape',
    'speaker1_image': 'Speaker1_Image_Shape',
    'speaker2_image': 'Speaker2_Image_Shape',
    'ge_image': 'GE_Image_Shape'
}

TEXT_SHAPE_NAMES = {
    'meeting_mode': 'Meeting Mode',
    'meeting_time': 'Meeting Time',
    'venue': 'Venue'
}


ROLE_TITLE_SHAPE_NAMES = {
    'ge_title': 'TM/GE' 
}



def get_date_suffix(day):
    if 4 <= day <= 20 or 24 <= day <= 30:
        return "th"
    return {1: "st", 2: "nd", 3: "rd"}.get(day % 10, "th")

def parse_date_from_fields(day_str: str, date_str: str, month_str: str, year_str: str) -> Optional[datetime]:
    """Attempts to reconstruct a date from individual fields."""
    try:
        
        
        date_num_str = ''.join(filter(str.isdigit, date_str))
        if not date_num_str: return None

        date_str_for_parsing = f"{date_num_str} {month_str} {year_str}"
        
        return datetime.strptime(date_str_for_parsing, "%d %B %Y")
    except ValueError:
        return None 

def find_image_file_path(name: str) -> Optional[str]:
    """Find an image file for a person in the server's images folder"""
    if not name:
        return None

    print(f"SERVER: Looking for image for: '{name}' in folder: {IMAGES_FOLDER}")
    extensions = ['.jpg', '.jpeg', '.png', '.gif', '.webp']
    for ext in extensions:
        
        image_path = os.path.join(IMAGES_FOLDER, f"{name}{ext}")
        if os.path.exists(image_path):
            print(f"SERVER: Found image at: {image_path}")
            return image_path

        
        image_path = os.path.join(IMAGES_FOLDER, f"{name.lower()}{ext}")
        if os.path.exists(image_path):
            print(f"SERVER: Found image at: {image_path}")
            return image_path

        
        no_spaces_name = name.replace(" ", "")
        image_path = os.path.join(IMAGES_FOLDER, f"{no_spaces_name}{ext}")
        if os.path.exists(image_path):
            print(f"SERVER: Found image at: {image_path}")
            return image_path
        image_path = os.path.join(IMAGES_FOLDER, f"{no_spaces_name.lower()}{ext}")
        if os.path.exists(image_path):
            print(f"SERVER: Found image at: {image_path}")
            return image_path

    print(f"SERVER: No image found for: '{name}'")
    return None

def update_text_preserving_format(shape, new_text):
    """Updates text in a shape while trying to preserve the formatting of the first run."""
    
    
    try:
        if not hasattr(shape, "text_frame") or not shape.text_frame:
            print(f"Warning: Shape {shape.shape_id} has no text_frame.")
            if hasattr(shape, "text"): 
                 shape.text = new_text
            return

        text_frame = shape.text_frame
        
        if len(text_frame.paragraphs) == 0:
            print(f"Debug: Text frame for shape {shape.shape_id} has 0 paragraphs. Setting text directly.")
            text_frame.text = new_text
            
            if text_frame.paragraphs:
                 text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER 
            return

        first_para = text_frame.paragraphs[0]
        alignment = first_para.alignment 
        level = first_para.level         

        
        first_run_format = {}
        if first_para.runs:
            first_run = first_para.runs[0]
            font = first_run.font
            first_run_format = {
                'name': font.name,
                'size': font.size,
                'bold': font.bold,
                'italic': font.italic,
                'underline': font.underline,
                'color_rgb': font.color.rgb if font.color and hasattr(font.color, 'rgb') else None
            }
            print(f"Debug: Preserving format from run: {first_run_format}")


        
        
        while len(text_frame.paragraphs) > 1:
            p_idx = len(text_frame.paragraphs) - 1
            p = text_frame.paragraphs[p_idx]
            text_frame._element.remove(p._element)


        first_para.clear() 

        
        run = first_para.add_run()
        run.text = new_text

        
        font = run.font
        if first_run_format.get('name') is not None: font.name = first_run_format['name']
        if first_run_format.get('size') is not None: font.size = first_run_format['size']
        if first_run_format.get('bold') is not None: font.bold = first_run_format['bold']
        if first_run_format.get('italic') is not None: font.italic = first_run_format['italic']
        if first_run_format.get('underline') is not None: font.underline = first_run_format['underline']
        if first_run_format.get('color_rgb') is not None: font.color.rgb = first_run_format['color_rgb']

        
        first_para.alignment = alignment
        first_para.level = level

        print(f"Debug: Successfully updated text for shape {shape.shape_id} preserving format.")

    except Exception as e:
        print(f"Error in update_text_preserving_format for shape {getattr(shape, 'shape_id', 'N/A')}: {str(e)}. Falling back.")
        
        try:
            if hasattr(shape, "text_frame") and shape.text_frame:
                shape.text_frame.text = new_text
            elif hasattr(shape, "text"):
                 shape.text = new_text
            else:
                 print(f"Error: Cannot set text for shape {getattr(shape, 'shape_id', 'N/A')}")
        except Exception as fallback_e:
            print(f"Fallback text setting failed for shape {getattr(shape, 'shape_id', 'N/A')}: {fallback_e}")




def analyze_presentation_data(ppt_stream: io.BytesIO) -> Dict[str, Any]:
    """Analyzes the presentation content from a byte stream."""
    try:
        prs = Presentation(ppt_stream)
    except PackageNotFoundError:
         raise ValueError("Invalid PowerPoint file format or corrupted file.")
    except Exception as e:
        raise ValueError(f"Could not open presentation: {e}")


    analysis = {
        'text_content_found': {key: [] for key in DEFAULT_SEARCH_TEXTS.keys()},
        'image_placeholders_found': {key: None for key in IMAGE_SHAPE_NAMES.keys()},
        'text_placeholders_found': {key: None for key in TEXT_SHAPE_NAMES.keys()},
        'role_title_placeholders_found': {key: None for key in ROLE_TITLE_SHAPE_NAMES.keys()},
        'initial_values': { 
             'theme': DEFAULT_SEARCH_TEXTS['theme'],
             'tmod': DEFAULT_SEARCH_TEXTS['tmod'].replace('TM ','').replace('TMOD NAME',''), 
             'ge': DEFAULT_SEARCH_TEXTS['ge'].replace('GE Name',''),
             'speaker1': DEFAULT_SEARCH_TEXTS['speaker1'].replace('TM Person1',''),
             'speaker2': DEFAULT_SEARCH_TEXTS['speaker2'].replace('TM Person2',''),
             'meeting_mode': '',
             'meeting_time': '',
             'venue': '',
             
             'day': '',
             'date': '',
             'month': '',
             'year': '',
        },
        'debug_log': [],
        'warnings': [],
        'status': "Analysis Started"
    }
    debug_log = analysis['debug_log']
    found_image_keys = set()
    found_text_keys = set()
    found_role_title_keys = set()
    extracted_date_fields = {'day': None, 'date': None, 'month': None, 'year': None}

    for slide_idx, slide in enumerate(prs.slides):
        for shape_idx, shape in enumerate(slide.shapes):
            shape_name = shape.name if hasattr(shape, 'name') and shape.name else None
            shape_text = None
            
            if hasattr(shape, "text_frame") and shape.text_frame and shape.text_frame.text.strip():
                 try:
                     shape_text = shape.text_frame.text.strip()
                 except Exception as text_read_err:
                      debug_log.append(f"Warning: Could not read text from shape idx {shape_idx} on slide {slide_idx+1}. Error: {text_read_err}")


            
            if shape_text:
                for key, search_text in DEFAULT_SEARCH_TEXTS.items():
                    if search_text.lower() in shape_text.lower():
                        debug_log.append(f"Found text content match for '{key}': '{shape_text}' on slide {slide_idx+1}, shape idx {shape_idx} (Name: '{shape_name}')")
                        analysis['text_content_found'][key].append({
                            'slide_idx': slide_idx,
                            'shape_idx': shape_idx, 
                            'shape_id': shape.shape_id, 
                            'current_text': shape_text
                        })
                        
                        if key in extracted_date_fields and extracted_date_fields[key] is None:
                             extracted_date_fields[key] = shape_text


            
            if shape_name:
                
                for key, expected_name in IMAGE_SHAPE_NAMES.items():
                    if expected_name.lower() == shape_name.lower():
                        debug_log.append(f"Found image placeholder for '{key}': Name '{shape_name}' on slide {slide_idx+1}, shape idx {shape_idx} (ID: {shape.shape_id})")
                        if key not in found_image_keys:
                            analysis['image_placeholders_found'][key] = {'slide_idx': slide_idx, 'shape_id': shape.shape_id}
                            found_image_keys.add(key)
                        break 

                
                for key, expected_name in TEXT_SHAPE_NAMES.items():
                    if expected_name.lower() == shape_name.lower():
                        current_text = shape_text if shape_text else ""
                        debug_log.append(f"Found text placeholder for '{key}': Name '{shape_name}' (Text: '{current_text}') on slide {slide_idx+1}, shape idx {shape_idx} (ID: {shape.shape_id})")
                        if key not in found_text_keys:
                            analysis['text_placeholders_found'][key] = {'slide_idx': slide_idx, 'shape_id': shape.shape_id, 'current_text': current_text}
                            found_text_keys.add(key)
                            
                            if key == 'meeting_mode': analysis['initial_values']['meeting_mode'] = current_text
                            elif key == 'meeting_time': analysis['initial_values']['meeting_time'] = current_text
                            elif key == 'venue': analysis['initial_values']['venue'] = current_text
                        break 

                
                for key, expected_name in ROLE_TITLE_SHAPE_NAMES.items():
                     if expected_name.lower() == shape_name.lower():
                        current_text = shape_text if shape_text else ""
                        debug_log.append(f"Found role title placeholder for '{key}': Name '{shape_name}' (Text: '{current_text}') on slide {slide_idx+1}, shape idx {shape_idx} (ID: {shape.shape_id})")
                        if key not in found_role_title_keys:
                            analysis['role_title_placeholders_found'][key] = {'slide_idx': slide_idx, 'shape_id': shape.shape_id, 'current_text': current_text}
                            found_role_title_keys.add(key)
                            
                        break 

    

    
    if analysis['text_content_found']['theme']:
        analysis['initial_values']['theme'] = analysis['text_content_found']['theme'][0]['current_text']

    if analysis['text_content_found']['tmod']:
        tmod_text = analysis['text_content_found']['tmod'][0]['current_text']
        analysis['initial_values']['tmod'] = tmod_text.replace(DEFAULT_SEARCH_TEXTS['tmod'], '').strip() 
        if not analysis['initial_values']['tmod']: 
             analysis['initial_values']['tmod'] = tmod_text[3:] if tmod_text.startswith('TM ') else tmod_text


    if analysis['text_content_found']['ge']:
        ge_text = analysis['text_content_found']['ge'][0]['current_text']
        analysis['initial_values']['ge'] = ge_text.replace(DEFAULT_SEARCH_TEXTS['ge'], '').strip()
        if not analysis['initial_values']['ge']:
             analysis['initial_values']['ge'] = ge_text

    if analysis['text_content_found']['speaker1']:
        sp1_text = analysis['text_content_found']['speaker1'][0]['current_text']
        analysis['initial_values']['speaker1'] = sp1_text.replace(DEFAULT_SEARCH_TEXTS['speaker1'], '').strip()
        if not analysis['initial_values']['speaker1']:
            analysis['initial_values']['speaker1'] = sp1_text[3:] if sp1_text.startswith('TM ') else sp1_text

    if analysis['text_content_found']['speaker2']:
        sp2_text = analysis['text_content_found']['speaker2'][0]['current_text']
        analysis['initial_values']['speaker2'] = sp2_text.replace(DEFAULT_SEARCH_TEXTS['speaker2'], '').strip()
        if not analysis['initial_values']['speaker2']:
            analysis['initial_values']['speaker2'] = sp2_text[3:] if sp2_text.startswith('TM ') else sp2_text

    
    current_date = datetime.now().date()
    analysis['initial_values']['day'] = extracted_date_fields.get('day') or current_date.strftime("%a").upper()
    analysis['initial_values']['date'] = extracted_date_fields.get('date') or (current_date.strftime("%d") + get_date_suffix(current_date.day))
    analysis['initial_values']['month'] = extracted_date_fields.get('month') or current_date.strftime("%B").upper()
    analysis['initial_values']['year'] = extracted_date_fields.get('year') or current_date.strftime("%Y")


    
    num_text_content_found = sum(len(shapes) for shapes in analysis['text_content_found'].values())
    num_images_found = len(found_image_keys)
    num_text_name_found = len(found_text_keys)
    num_role_title_found = len(found_role_title_keys)

    status_parts = [
        f"Analysis Complete.",
        f"{num_text_content_found} role/theme/date texts potentially found by content.",
        f"{num_images_found}/{len(IMAGE_SHAPE_NAMES)} image placeholders found by name.",
        f"{num_text_name_found}/{len(TEXT_SHAPE_NAMES)} detail text placeholders found by name.",
        f"{num_role_title_found}/{len(ROLE_TITLE_SHAPE_NAMES)} role title placeholders found by name."
    ]
    analysis['status'] = " ".join(status_parts)

    if num_images_found < len(IMAGE_SHAPE_NAMES):
        missing = [k for k,v in analysis['image_placeholders_found'].items() if v is None]
        analysis['warnings'].append(f"Could not find image placeholders by Shape Name for: {', '.join(missing)}. Expected names: {[IMAGE_SHAPE_NAMES[k] for k in missing]}")
    if num_text_name_found < len(TEXT_SHAPE_NAMES):
        missing = [k for k,v in analysis['text_placeholders_found'].items() if v is None]
        analysis['warnings'].append(f"Could not find detail text placeholders by Shape Name for: {', '.join(missing)}. Expected names: {[TEXT_SHAPE_NAMES[k] for k in missing]}")
    if num_role_title_found < len(ROLE_TITLE_SHAPE_NAMES):
        missing = [k for k,v in analysis['role_title_placeholders_found'].items() if v is None]
        analysis['warnings'].append(f"Could not find role title placeholders by Shape Name for: {', '.join(missing)}. Expected names: {[ROLE_TITLE_SHAPE_NAMES[k] for k in missing]}")

    
    
    return analysis


def update_presentation_data(ppt_stream: io.BytesIO, update_data: Dict[str, Any], analysis_results: Dict[str, Any]) -> Tuple[bytes, List[str]]:
    """Updates the presentation based on provided data and analysis results."""
    try:
        prs = Presentation(ppt_stream)
    except Exception as e:
        raise ValueError(f"Could not open presentation for updating: {e}")

    debug_log = []
    replacements = {
        'theme': update_data['theme'],
        'day': update_data['day'],
        'date': update_data['date'],
        'month': update_data['month'],
        'year': update_data['year'],
        'tmod': f"TM {update_data['tmod']}" if update_data.get('tmod') else "", 
        'ge': update_data['ge'], 
        'speaker1': f"TM {update_data['speaker1']}" if update_data.get('speaker1') else "",
        'speaker2': f"TM {update_data['speaker2']}" if update_data.get('speaker2') else "",
        'meeting_mode': update_data['meeting_mode'],
        'meeting_time': update_data['meeting_time'],
        'venue': update_data['venue'],
    }

    
    debug_log.append("--- Updating Text (Roles/Theme/Date by Content Match) ---")
    content_matches = analysis_results.get('text_content_found', {})
    keys_to_update_by_content = ['theme', 'day', 'date', 'month', 'year', 'tmod', 'ge', 'speaker1', 'speaker2']

    for key in keys_to_update_by_content:
        if key not in replacements: continue 
        new_text = replacements[key]
        shapes_found = content_matches.get(key, [])

        if not shapes_found:
             debug_log.append(f"No content match shapes found for '{key}' during analysis.")
             continue

        for shape_info in shapes_found:
            try:
                slide = prs.slides[shape_info['slide_idx']]
                
                shape = next((s for s in slide.shapes if s.shape_id == shape_info['shape_id']), None)

                if shape:
                    original_search_text = DEFAULT_SEARCH_TEXTS.get(key)
                    current_text_in_shape = shape.text if hasattr(shape, 'text') else ''

                    
                    if key in ['tmod', 'ge', 'speaker1', 'speaker2'] and original_search_text and original_search_text.lower() in current_text_in_shape.lower():
                         
                         
                         
                         import re
                         try:
                             
                             pattern = re.compile(re.escape(original_search_text), re.IGNORECASE)
                             updated_text = pattern.sub(new_text, current_text_in_shape)
                             if updated_text != current_text_in_shape: 
                                 update_text_preserving_format(shape, updated_text)
                                 debug_log.append(f"Updated '{key}' on slide {shape_info['slide_idx']+1} (Shape ID: {shape.shape_id}) via partial text replace.")
                             else:
                                 
                                 update_text_preserving_format(shape, new_text)
                                 debug_log.append(f"Updated '{key}' on slide {shape_info['slide_idx']+1} (Shape ID: {shape.shape_id}) via full replace (partial failed).")
                         except Exception as replace_err:
                              debug_log.append(f"Warning: Regex replace failed for '{key}' (Shape ID: {shape.shape_id}). Error: {replace_err}. Falling back to full replace.")
                              update_text_preserving_format(shape, new_text)
                              debug_log.append(f"Updated '{key}' on slide {shape_info['slide_idx']+1} (Shape ID: {shape.shape_id}) via full replace (exception fallback).")

                    else:
                        
                        update_text_preserving_format(shape, new_text)
                        debug_log.append(f"Updated '{key}' on slide {shape_info['slide_idx']+1} (Shape ID: {shape.shape_id}) via full text replace.")
                else:
                    debug_log.append(f"Warning: Shape ID {shape_info['shape_id']} not found on slide {shape_info['slide_idx']+1} when updating '{key}'.")
            except IndexError:
                 debug_log.append(f"Warning: Slide index {shape_info['slide_idx']} out of range when updating '{key}'.")
            except Exception as text_e:
                 debug_log.append(f"Warning: Error updating text content for '{key}' (Shape ID: {shape_info.get('shape_id','N/A')}): {str(text_e)}")


    
    debug_log.append("--- Updating Text (Details by Shape Name) ---")
    text_placeholders = analysis_results.get('text_placeholders_found', {})
    for key, placeholder_info in text_placeholders.items():
        if placeholder_info and key in replacements:
            new_text = replacements[key]
            slide_idx = placeholder_info['slide_idx']
            shape_id = placeholder_info['shape_id']
            expected_shape_name = TEXT_SHAPE_NAMES.get(key)
            debug_log.append(f"Attempting to update text for '{key}' ('{new_text}') by finding shape named '{expected_shape_name}' (ID: {shape_id}) on slide {slide_idx+1}")
            try:
                if slide_idx >= len(prs.slides):
                    debug_log.append(f"Slide index {slide_idx} out of range.")
                    continue
                slide = prs.slides[slide_idx]
                target_shape = next((s for s in slide.shapes if s.shape_id == shape_id), None)

                if target_shape and hasattr(target_shape, 'name') and target_shape.name == expected_shape_name:
                    update_text_preserving_format(target_shape, new_text)
                    debug_log.append(f"Successfully updated text for '{key}' by name (ID: {shape_id}).")
                elif target_shape:
                     debug_log.append(f"Warning: Found shape ID {shape_id} but name '{target_shape.name}' != expected '{expected_shape_name}' for '{key}'. Skipping update.")
                else:
                     debug_log.append(f"Could not find shape ID {shape_id} on slide {slide_idx+1} during text update for '{key}'.")
            except Exception as e:
                debug_log.append(f"Error updating text placeholder '{key}' (ID: {shape_id}): {str(e)}")
        elif key in replacements:
             debug_log.append(f"Text placeholder '{key}' not found during analysis. Skipping update by name.")


    
    debug_log.append("--- Updating Images (by Shape Name) ---")
    image_placeholders = analysis_results.get('image_placeholders_found', {})
    
    image_data_map = {
        'tmod_image':     {'name': update_data.get('tmod'),     'filename': update_data.get('tmod_image')},
        'speaker1_image': {'name': update_data.get('speaker1'), 'filename': update_data.get('speaker1_image')},
        'speaker2_image': {'name': update_data.get('speaker2'), 'filename': update_data.get('speaker2_image')},
        'ge_image':       {'name': update_data.get('ge'),       'filename': update_data.get('ge_image')}
    }

    for placeholder_key, data in image_data_map.items():
        placeholder_info = image_placeholders.get(placeholder_key)
        person_name = data['name']
        specific_filename = data['filename']

        if not placeholder_info:
            debug_log.append(f"Image placeholder '{placeholder_key}' not found in analysis. Skipping image update.")
            continue
        
        if not person_name and not specific_filename:
            debug_log.append(f"No person name or specific filename provided for '{placeholder_key}'. Skipping image update.")
            continue

        slide_idx = placeholder_info['slide_idx']
        shape_id = placeholder_info['shape_id']
        expected_shape_name = IMAGE_SHAPE_NAMES.get(placeholder_key)
        try:
            if slide_idx >= len(prs.slides):
                debug_log.append(f"Warning: Slide index {slide_idx} out of range for image placeholder '{placeholder_key}'.")
                continue
            slide = prs.slides[slide_idx]
            target_shape = next((s for s in slide.shapes if s.shape_id == shape_id), None)

            if target_shape is None:
                debug_log.append(f"Could not find shape ID {shape_id} for image placeholder '{placeholder_key}'.")
                continue
            if not hasattr(target_shape, 'name') or target_shape.name != expected_shape_name:
                 debug_log.append(f"Warning: Found shape ID {shape_id} but name '{getattr(target_shape,'name','N/A')}' != expected '{expected_shape_name}' for '{placeholder_key}'. Skipping.")
                 continue

            
            image_path = None
            if specific_filename:
                
                potential_path = os.path.join(IMAGES_FOLDER, specific_filename)
                
                if os.path.abspath(potential_path).startswith(os.path.abspath(IMAGES_FOLDER)) and os.path.exists(potential_path):
                    image_path = potential_path
                    debug_log.append(f"Using specific image file: {image_path} for {placeholder_key}")
                else:
                     debug_log.append(f"Warning: Specific filename '{specific_filename}' provided for '{placeholder_key}' but not found or invalid path. Falling back to name search.")
            
            
            if not image_path and person_name:
                debug_log.append(f"Searching for image by name: '{person_name}' for {placeholder_key}")
                image_path = find_image_file_path(person_name) 

            
            if image_path:
                left, top, width, height = target_shape.left, target_shape.top, target_shape.width, target_shape.height
                debug_log.append(f"Found image {image_path}. Replacing placeholder {placeholder_key} (ID: {shape_id}).")
                try:
                    old_element = target_shape._element
                    parent_element = old_element.getparent()
                    if parent_element is None: raise ValueError("Could not get parent element.")
                    parent_element.remove(old_element)
                    new_pic = slide.shapes.add_picture(image_path, left, top, width, height)
                    new_pic.name = expected_shape_name 
                    debug_log.append(f"Successfully replaced image for '{placeholder_key}' (New ID: {new_pic.shape_id}) using {image_path}")
                except Exception as replace_e:
                    debug_log.append(f"ERROR during image replacement for '{placeholder_key}' (ID: {shape_id}): {str(replace_e)}")
            else:
                debug_log.append(f"No image file found for placeholder '{placeholder_key}' (Name: '{person_name}', Filename: '{specific_filename}'). Skipping.")

        except Exception as e:
            debug_log.append(f"Error processing image placeholder '{placeholder_key}' (ID: {shape_id}): {str(e)}")


    debug_log.append("--- Updating Role Title (by Shape Name) ---")
    role_title_placeholders = analysis_results.get('role_title_placeholders_found', {})
    for key, placeholder_info in role_title_placeholders.items():
        
        if placeholder_info and key == 'ge_title': 
            new_text = update_data['ge_title']
            slide_idx = placeholder_info['slide_idx']
            shape_id = placeholder_info['shape_id']
            expected_shape_name = ROLE_TITLE_SHAPE_NAMES.get(key, "")
            try:
                if slide_idx >= len(prs.slides):
                    debug_log.append(f"Warning: Slide index {slide_idx} out of range for role title placeholder '{key}'.")
                    continue
                slide = prs.slides[slide_idx]
                target_shape = next((s for s in slide.shapes if s.shape_id == shape_id), None)

                if target_shape and hasattr(target_shape, 'name') and target_shape.name.lower() == expected_shape_name.lower():
                    
                    if new_text.lower() == 'speaker 3':
                        sp2_matches = analysis_results.get('text_content_found', {}).get('speaker2', [])
                        if sp2_matches:
                            sp2_info = sp2_matches[0]
                            sp2_slide = prs.slides[sp2_info['slide_idx']]
                            sp2_shape = next((s for s in sp2_slide.shapes if s.shape_id == sp2_info['shape_id']), None)
                            if sp2_shape and hasattr(sp2_shape, 'text_frame') and sp2_shape.text_frame.paragraphs:
                                first_sp2_para = sp2_shape.text_frame.paragraphs[0]
                                
                                sp2_font = first_sp2_para.runs[0].font if first_sp2_para.runs else None
                                
                                tf = target_shape.text_frame
                                while len(tf.paragraphs) > 1:
                                    p = tf.paragraphs[-1]; tf._element.remove(p._element)
                                tf.paragraphs[0].clear()
                                run = tf.paragraphs[0].add_run()
                                run.text = new_text 
                                if sp2_font:
                                    f = run.font
                                    if sp2_font.name: f.name = sp2_font.name
                                    if sp2_font.size: f.size = sp2_font.size
                                    if sp2_font.bold is not None: f.bold = sp2_font.bold
                                    if sp2_font.italic is not None: f.italic = sp2_font.italic
                                    if sp2_font.underline is not None: f.underline = sp2_font.underline
                                    if sp2_font.color and hasattr(sp2_font.color, 'rgb'): f.color.rgb = sp2_font.color.rgb
                                tf.paragraphs[0].alignment = first_sp2_para.alignment
                                tf.paragraphs[0].level = first_sp2_para.level
                                debug_log.append(f"Updated '{key}' using Speaker2 formatting (ID: {shape_id}).")
                            else:
                                update_text_preserving_format(target_shape, new_text)
                                debug_log.append(f"Updated '{key}' via fallback full text replace.")
                        else:
                            update_text_preserving_format(target_shape, new_text)
                            debug_log.append(f"Updated '{key}' via fallback full text replace.")
                    else:
                        
                        update_text_preserving_format(target_shape, new_text)
                        debug_log.append(f"Successfully updated role title for '{key}' by name (ID: {shape_id}).")
                else:
                    debug_log.append(f"Warning: Found shape ID {shape_id} but name '{target_shape.name}' doesn't match expected '{expected_shape_name}' for role title '{key}'. Skipping update.")
            except Exception as e:
                debug_log.append(f"Error updating role title placeholder '{key}' (ID: {shape_id}): {str(e)}")
        elif key == 'ge_title': 
             debug_log.append(f"Role title placeholder '{key}' not found during analysis. Skipping update by name.")

    
    output_stream = io.BytesIO()
    prs.save(output_stream)
    output_stream.seek(0)

    return output_stream.getvalue(), debug_log



app = FastAPI(title="Slide Style Sync Backend")


origins = [
    "http://localhost:3000",  
    "http://127.0.0.1:3000",
    "http://localhost:5173",  
    "http://127.0.0.1:5173", 
    "http://localhost:8080",  
    "http://127.0.0.1:8080", 
    
]

app.add_middleware(
    CORSMiddleware,
    allow_origins=origins,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


class AnalysisResult(BaseModel):
    fileName: str
    analysisDetails: Dict[str, Any] 
    status: str
    warnings: Optional[List[str]] = None
    debugLog: Optional[List[str]] = None
    fileId: str 

class UpdateRequestData(BaseModel):
    fileId: str 
    theme: str
    
    day: str
    date: str
    month: str
    year: str
    
    tmod: str = ""
    ge: str = ""
    speaker1: str = ""
    speaker2: str = ""
    
    
    
    meeting_mode: str = Field(alias="meetingMode", default="")
    meeting_time: str = Field(alias="meetingTime", default="")
    venue: str = ""
    
    tmod_image: Optional[str] = None
    ge_image: Optional[str] = None
    speaker1_image: Optional[str] = None
    speaker2_image: Optional[str] = None
    
    ge_title: str = "General Evaluator" 

class ImageListResponse(BaseModel):
    images: List[str]


class Profile(BaseModel):
    name: str
    image_filename: str

class ProfileListResponse(BaseModel):
    profiles: List[Profile]

class ProfileUpdateResponse(BaseModel):
    message: str
    profile: Profile

class ProfileDeleteResponse(BaseModel):
    message: str
    name: str




uploaded_file_data: Dict[str, Dict[str, Any]] = {}



@app.post("/upload_analyze/", response_model=AnalysisResult)
async def upload_and_analyze_ppt(file: UploadFile = File(...)):
    """
    Receives a PPTX file, stores it temporarily, analyzes it,
    and returns the analysis results along with a fileId.
    """
    if not file.filename or not file.filename.lower().endswith(".pptx"):
        raise HTTPException(status_code=400, detail="Invalid file type. Please upload a .pptx file.")

    
    file_id = str(uuid.uuid4())
    file_content = None
    analysis_data = {}
    status = "Error"
    warnings_list = []
    debug_log = []

    try:
        debug_log.append(f"Received file: {file.filename}, size: {file.size}, type: {file.content_type}")
        file_content = await file.read()
        if not file_content:
             raise ValueError("Uploaded file is empty.")
        debug_log.append(f"File read into memory ({len(file_content)} bytes). Starting analysis.")

        file_stream = io.BytesIO(file_content)
        analysis_data = analyze_presentation_data(file_stream)
        debug_log.extend(analysis_data.get('debug_log', []))
        warnings_list = analysis_data.get('warnings', [])
        status = analysis_data.get('status', "Analysis complete, but status missing.")
        debug_log.append(f"Analysis finished. Status: {status}")

        
        uploaded_file_data[file_id] = {
            "filename": file.filename,
            "content": file_content,
            "analysis": analysis_data
        }
        debug_log.append(f"Stored file content and analysis under ID: {file_id}")

        
        
        client_analysis_details = analysis_data.copy()
        if 'debug_log' in client_analysis_details: del client_analysis_details['debug_log']
        
        


        return AnalysisResult(
            fileName=file.filename,
            analysisDetails=client_analysis_details, 
            status=status,
            warnings=warnings_list,
            
            fileId=file_id
        )

    except ValueError as ve:
         debug_log.append(f"Value Error during processing: {ve}")
         raise HTTPException(status_code=400, detail=str(ve))
    except Exception as e:
        debug_log.append(f"Unexpected Error during upload/analysis: {e}")
        import traceback
        traceback.print_exc() 
        
        if file_id in uploaded_file_data:
            del uploaded_file_data[file_id]
        raise HTTPException(status_code=500, detail=f"Internal server error during analysis: {str(e)}")
    finally:
        await file.close()
        print(f"--- Upload/Analyze Log for {file.filename} (ID: {file_id}) ---")
        for entry in debug_log: print(entry)
        print("--- End Log ---")


@app.post("/update_presentation/")
async def update_ppt_and_get_image(update_request: UpdateRequestData):
    """
    Updates the presentation data, saves it temporarily, converts the first slide
    to PNG using pptx_tools, crops it, and returns the cropped image.
    """
    if not save_pptx_as_png or not PILImage:
        raise HTTPException(status_code=500, detail="Server is missing required image processing libraries (Pillow, python-pptx-tools, or dependencies). Please check server logs.")

    file_id = update_request.fileId
    debug_log = []
    print(f"Received update request for fileId: {file_id}")

    if file_id not in uploaded_file_data:
        raise HTTPException(status_code=404, detail=f"File session '{file_id}' not found or expired.")

    stored_data = uploaded_file_data[file_id]
    original_filename = stored_data["filename"]
    original_content = stored_data["content"]
    analysis_results = stored_data["analysis"]

    
    TARGET_WIDTH = 1920 
    TARGET_HEIGHT = 1080 
    TARGET_SLIDE_INDEX = 0 
    expected_png_filename = f"slide_{TARGET_SLIDE_INDEX+1:03d}.png" 
    alt_expected_png_filename = f"slide{TARGET_SLIDE_INDEX+1}.png"

    temp_pptx_path = None
    temp_png_dir = None

    try:
        
        debug_log.append(f"Loading '{original_filename}' for updates...")
        original_stream = io.BytesIO(original_content)
        prs_for_update = Presentation(original_stream)
        debug_log.append(f"Loaded presentation with {len(prs_for_update.slides)} slides.")

        
        debug_log.append("--- Applying updates to presentation object ---")
        
        
        replacements = {
            'theme': update_request.theme,
            'day': update_request.day, 'date': update_request.date, 'month': update_request.month, 'year': update_request.year,
            'tmod': f"TM {update_request.tmod}" if update_request.tmod else "", 
            'ge': update_request.ge,
            'speaker1': f"TM {update_request.speaker1}" if update_request.speaker1 else "", 
            'speaker2': f"TM {update_request.speaker2}" if update_request.speaker2 else "",
            'meeting_mode': update_request.meeting_mode,
            'meeting_time': update_request.meeting_time,
            'venue': update_request.venue,
        }
        debug_log.append(f"Replacements prepared: {replacements}")

        
        content_matches = analysis_results.get('text_content_found', {})
        keys_to_update_by_content = [
            'theme', 'day', 'date', 'month', 'year', 
            'tmod', 'ge', 'speaker1', 'speaker2'
            
        ]
        for key in keys_to_update_by_content:
            if key not in replacements: continue
            new_text = replacements[key]
            shapes_found = content_matches.get(key, [])
            if not shapes_found:
                debug_log.append(f"No content match shapes found for '{key}' during analysis.")
                continue
            
            for shape_info in shapes_found:
                try:
                    slide_idx = shape_info['slide_idx']
                    shape_id = shape_info['shape_id']
                    if slide_idx >= len(prs_for_update.slides):
                        debug_log.append(f"Warning: Slide index {slide_idx} out of range when updating '{key}'.")
                        continue
                    
                    slide = prs_for_update.slides[slide_idx]
                    
                    shape = next((s for s in slide.shapes if s.shape_id == shape_id), None)

                    if shape:
                        original_search_text = DEFAULT_SEARCH_TEXTS.get(key)
                        current_text_in_shape = shape.text if hasattr(shape, 'text') else ''

                        
                        if key in ['tmod', 'ge', 'speaker1', 'speaker2'] and original_search_text and original_search_text.lower() in current_text_in_shape.lower():
                            
                            
                            
                            import re
                            try:
                                pattern = re.compile(re.escape(original_search_text), re.IGNORECASE)
                                updated_text = pattern.sub(new_text, current_text_in_shape)
                                if updated_text != current_text_in_shape:
                                    update_text_preserving_format(shape, updated_text)
                                    debug_log.append(f"Updated '{key}' (ID: {shape_id}) via partial replace.")
                                else:
                                    update_text_preserving_format(shape, new_text)
                                    debug_log.append(f"Updated '{key}' (ID: {shape_id}) via full replace (partial failed).")
                            except Exception as replace_err:
                                debug_log.append(f"Warning: Regex replace failed for '{key}' (ID: {shape_id}). Error: {replace_err}. Falling back to full replace.")
                                update_text_preserving_format(shape, new_text)
                                debug_log.append(f"Updated '{key}' (ID: {shape_id}) via full replace (exception fallback).")

                        else:
                            
                            update_text_preserving_format(shape, new_text)
                            debug_log.append(f"Updated '{key}' (ID: {shape_id}) via full text replace.")
                    else:
                        debug_log.append(f"Warning: Shape ID {shape_id} not found on slide {slide_idx} when updating '{key}'.")
                except Exception as text_e:
                     debug_log.append(f"Warning: Error updating text content for '{key}' (Shape ID: {shape_info.get('shape_id','N/A')}): {str(text_e)}")


        
        text_placeholders = analysis_results.get('text_placeholders_found', {})
        for key, placeholder_info in text_placeholders.items():
            if placeholder_info and key in replacements:
                new_text = replacements[key]
                slide_idx = placeholder_info['slide_idx']
                shape_id = placeholder_info['shape_id']
                expected_shape_name = TEXT_SHAPE_NAMES.get(key)
                try:
                    if slide_idx >= len(prs_for_update.slides):
                        debug_log.append(f"Warning: Slide index {slide_idx} out of range for text placeholder '{key}'.")
                        continue
                    slide = prs_for_update.slides[slide_idx]
                    target_shape = next((s for s in slide.shapes if s.shape_id == shape_id), None)

                    if target_shape and hasattr(target_shape, 'name') and target_shape.name == expected_shape_name:
                        update_text_preserving_format(target_shape, new_text)
                        debug_log.append(f"Updated text placeholder '{key}' (ID: {shape_id}) by name.")
                    elif target_shape:
                        debug_log.append(f"Warning: Found shape ID {shape_id} but name '{target_shape.name}' != expected '{expected_shape_name}' for '{key}'. Skipping.")
                    else:
                         debug_log.append(f"Could not find shape ID {shape_id} for text placeholder '{key}'.")
                except Exception as e:
                    debug_log.append(f"Error updating text placeholder '{key}' (ID: {shape_id}): {str(e)}")
            elif key in replacements:
                 debug_log.append(f"Text placeholder '{key}' not found during analysis. Skipping update by name.")


        
        image_placeholders = analysis_results.get('image_placeholders_found', {})
        
        image_data_map = {
            'tmod_image':     {'name': update_request.tmod,     'filename': update_request.tmod_image},
            'speaker1_image': {'name': update_request.speaker1, 'filename': update_request.speaker1_image},
            'speaker2_image': {'name': update_request.speaker2, 'filename': update_request.speaker2_image},
            'ge_image':       {'name': update_request.ge,       'filename': update_request.ge_image}
        }

        for placeholder_key, data in image_data_map.items():
            placeholder_info = image_placeholders.get(placeholder_key)
            person_name = data['name']
            specific_filename = data['filename']

            if not placeholder_info:
                debug_log.append(f"Image placeholder '{placeholder_key}' not found in analysis. Skipping image update.")
                continue
            
            if not person_name and not specific_filename:
                debug_log.append(f"No person name or specific filename provided for '{placeholder_key}'. Skipping image update.")
                continue

            slide_idx = placeholder_info['slide_idx']
            shape_id = placeholder_info['shape_id']
            expected_shape_name = IMAGE_SHAPE_NAMES.get(placeholder_key)
            try:
                if slide_idx >= len(prs_for_update.slides):
                    debug_log.append(f"Warning: Slide index {slide_idx} out of range for image placeholder '{placeholder_key}'.")
                    continue
                slide = prs_for_update.slides[slide_idx]
                target_shape = next((s for s in slide.shapes if s.shape_id == shape_id), None)

                if target_shape is None:
                    debug_log.append(f"Could not find shape ID {shape_id} for image placeholder '{placeholder_key}'.")
                    continue
                if not hasattr(target_shape, 'name') or target_shape.name != expected_shape_name:
                     debug_log.append(f"Warning: Found shape ID {shape_id} but name '{getattr(target_shape,'name','N/A')}' != expected '{expected_shape_name}' for '{placeholder_key}'. Skipping.")
                     continue

                
                image_path = None
                if specific_filename:
                    
                    potential_path = os.path.join(IMAGES_FOLDER, specific_filename)
                    
                    if os.path.abspath(potential_path).startswith(os.path.abspath(IMAGES_FOLDER)) and os.path.exists(potential_path):
                        image_path = potential_path
                        debug_log.append(f"Using specific image file: {image_path} for {placeholder_key}")
                    else:
                         debug_log.append(f"Warning: Specific filename '{specific_filename}' provided for '{placeholder_key}' but not found or invalid path. Falling back to name search.")
                
                
                if not image_path and person_name:
                    debug_log.append(f"Searching for image by name: '{person_name}' for {placeholder_key}")
                    image_path = find_image_file_path(person_name) 

                
                if image_path:
                    left, top, width, height = target_shape.left, target_shape.top, target_shape.width, target_shape.height
                    debug_log.append(f"Found image {image_path}. Replacing placeholder {placeholder_key} (ID: {shape_id}).")
                    try:
                        old_element = target_shape._element
                        parent_element = old_element.getparent()
                        if parent_element is None: raise ValueError("Could not get parent element.")
                        parent_element.remove(old_element)
                        new_pic = slide.shapes.add_picture(image_path, left, top, width, height)
                        new_pic.name = expected_shape_name 
                        debug_log.append(f"Successfully replaced image for '{placeholder_key}' (New ID: {new_pic.shape_id}) using {image_path}")
                    except Exception as replace_e:
                        debug_log.append(f"ERROR during image replacement for '{placeholder_key}' (ID: {shape_id}): {str(replace_e)}")
                else:
                    debug_log.append(f"No image file found for placeholder '{placeholder_key}' (Name: '{person_name}', Filename: '{specific_filename}'). Skipping.")

            except Exception as e:
                debug_log.append(f"Error processing image placeholder '{placeholder_key}' (ID: {shape_id}): {str(e)}")

        
        debug_log.append("--- Updating Role Title (by Shape Name) ---")
        role_title_placeholders = analysis_results.get('role_title_placeholders_found', {})
        for key, placeholder_info in role_title_placeholders.items():
            
            if placeholder_info and key == 'ge_title': 
                new_text = update_request.ge_title
                slide_idx = placeholder_info['slide_idx']
                shape_id = placeholder_info['shape_id']
                expected_shape_name = ROLE_TITLE_SHAPE_NAMES.get(key, "")
                try:
                    if slide_idx >= len(prs_for_update.slides):
                        debug_log.append(f"Warning: Slide index {slide_idx} out of range for role title placeholder '{key}'.")
                        continue
                    slide = prs_for_update.slides[slide_idx]
                    target_shape = next((s for s in slide.shapes if s.shape_id == shape_id), None)

                    if target_shape and hasattr(target_shape, 'name') and target_shape.name.lower() == expected_shape_name.lower():
                        
                        if new_text.lower() == 'speaker 3':
                            sp2_matches = analysis_results.get('text_content_found', {}).get('speaker2', [])
                            if sp2_matches:
                                sp2_info = sp2_matches[0]
                                sp2_slide = prs_for_update.slides[sp2_info['slide_idx']]
                                sp2_shape = next((s for s in sp2_slide.shapes if s.shape_id == sp2_info['shape_id']), None)
                                if sp2_shape and hasattr(sp2_shape, 'text_frame') and sp2_shape.text_frame.paragraphs:
                                    first_sp2_para = sp2_shape.text_frame.paragraphs[0]
                                    
                                    sp2_font = first_sp2_para.runs[0].font if first_sp2_para.runs else None
                                    
                                    tf = target_shape.text_frame
                                    while len(tf.paragraphs) > 1:
                                        p = tf.paragraphs[-1]; tf._element.remove(p._element)
                                    tf.paragraphs[0].clear()
                                    run = tf.paragraphs[0].add_run()
                                    run.text = new_text 
                                    if sp2_font:
                                        f = run.font
                                        if sp2_font.name: f.name = sp2_font.name
                                        if sp2_font.size: f.size = sp2_font.size
                                        if sp2_font.bold is not None: f.bold = sp2_font.bold
                                        if sp2_font.italic is not None: f.italic = sp2_font.italic
                                        if sp2_font.underline is not None: f.underline = sp2_font.underline
                                        if sp2_font.color and hasattr(sp2_font.color, 'rgb'): f.color.rgb = sp2_font.color.rgb
                                    tf.paragraphs[0].alignment = first_sp2_para.alignment
                                    tf.paragraphs[0].level = first_sp2_para.level
                                    debug_log.append(f"Updated '{key}' using Speaker2 formatting (ID: {shape_id}).")
                                else:
                                    update_text_preserving_format(target_shape, new_text)
                                    debug_log.append(f"Updated '{key}' via fallback full text replace.")
                            else:
                                update_text_preserving_format(target_shape, new_text)
                                debug_log.append(f"Updated '{key}' via fallback full text replace.")
                        else:
                            
                            update_text_preserving_format(target_shape, new_text)
                            debug_log.append(f"Successfully updated role title for '{key}' by name (ID: {shape_id}).")
                    else:
                        debug_log.append(f"Warning: Found shape ID {shape_id} but name '{target_shape.name}' doesn't match expected '{expected_shape_name}' for role title '{key}'. Skipping update.")
                except Exception as e:
                    debug_log.append(f"Error updating role title placeholder '{key}' (ID: {shape_id}): {str(e)}")
            elif key == 'ge_title': 
                 debug_log.append(f"Role title placeholder '{key}' not found during analysis. Skipping update by name.")

        
        with tempfile.NamedTemporaryFile(delete=False, suffix=".pptx", dir=TEMP_FOLDER) as temp_pptx:
            temp_pptx_path = temp_pptx.name
            prs_for_update.save(temp_pptx_path)
        debug_log.append(f"Saved updated presentation to temporary file: {temp_pptx_path}")

        
        temp_png_dir = tempfile.mkdtemp(dir=TEMP_FOLDER)
        debug_log.append(f"Created temporary PNG directory: {temp_png_dir}")

        
        debug_log.append(f"Calling save_pptx_as_png...")
        try:
            save_pptx_as_png(temp_png_dir, os.path.abspath(temp_pptx_path), overwrite_folder=True)
            debug_log.append(f"save_pptx_as_png completed.")
        except Exception as conversion_error:
             debug_log.append(f"ERROR during save_pptx_as_png: {conversion_error}")
             
             traceback.print_exc()
             raise RuntimeError(f"Failed to convert PPTX to PNG: {conversion_error}")

        
        target_png_path = os.path.join(temp_png_dir, expected_png_filename)
        alt_target_png_path = os.path.join(temp_png_dir, alt_expected_png_filename)
        actual_png_path = None
        if os.path.exists(target_png_path):
             actual_png_path = target_png_path
        elif os.path.exists(alt_target_png_path):
             actual_png_path = alt_target_png_path
        else:
            png_files = [f for f in os.listdir(temp_png_dir) if f.lower().endswith('.png')]
            if png_files:
                 png_files.sort()
                 debug_log.append(f"Warning: Expected PNG name not found. Found: {png_files}. Guessing first slide is: {png_files[0]}")
                 actual_png_path = os.path.join(temp_png_dir, png_files[0])
            else:
                 debug_log.append(f"ERROR: No PNG files found in output directory {temp_png_dir}")
                 raise RuntimeError(f"Conversion failed: No PNG file generated for slide {TARGET_SLIDE_INDEX}")
        debug_log.append(f"Loading target PNG: {actual_png_path}")
        img = PILImage.open(actual_png_path).convert("RGBA")

        
        debug_log.append(f"Cropping image to {TARGET_WIDTH}x{TARGET_HEIGHT}...")
        img_width, img_height = img.size
        left = max(0, int((img_width - TARGET_WIDTH) / 2))
        top = max(0, int((img_height - TARGET_HEIGHT) / 2))
        right = min(img_width, int(left + TARGET_WIDTH))
        bottom = min(img_height, int(top + TARGET_HEIGHT))
        if right <= left or bottom <= top:
             debug_log.append(f"Warning: Target crop dimensions invalid for image size. Cropping skipped.")
             cropped_img = img
        else:
             cropped_img = img.crop((left, top, right, bottom))
             debug_log.append("Cropping complete.")

        
        cropped_bytes_io = io.BytesIO()
        cropped_img.save(cropped_bytes_io, format='PNG')
        cropped_bytes_io.seek(0)
        image_bytes = cropped_bytes_io.getvalue()
        debug_log.append(f"Cropped image prepared ({len(image_bytes)} bytes).")

        
        image_filename = f"{os.path.splitext(original_filename)[0]}_slide_{TARGET_SLIDE_INDEX}.png"
        
        
        
        return Response(
            content=image_bytes,
            media_type="image/png",
            headers={ "Content-Disposition": f'attachment; filename="{image_filename}"' } 
        )

    except Exception as e:
        
        debug_log.append(f"General Error during update/image generation: {e}")
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"Internal server error: {str(e)}")

    finally:
        
        if temp_pptx_path and os.path.exists(temp_pptx_path):
            try: os.remove(temp_pptx_path)
            except OSError as e: debug_log.append(f"Warning: Failed to delete temp PPTX {temp_pptx_path}: {e}")
        if temp_png_dir and os.path.exists(temp_png_dir):
            try: shutil.rmtree(temp_png_dir)
            except OSError as e: debug_log.append(f"Warning: Failed to delete temp PNG dir {temp_png_dir}: {e}")
        print(f"--- Update/Generate Image Log for {original_filename} (ID: {file_id}) ---")
        for entry in debug_log: print(entry)
        print("--- End Log ---")


@app.post("/upload_image/")
async def upload_image(file: UploadFile = File(...)):
    """Uploads an image file to the server's image directory."""
    allowed_extensions = {".jpg", ".jpeg", ".png", ".gif", ".webp"}
    filename = file.filename
    extension = os.path.splitext(filename)[1].lower()

    if extension not in allowed_extensions:
        raise HTTPException(status_code=400, detail=f"Invalid image file type. Allowed: {', '.join(allowed_extensions)}")

    
    
    safe_filename = filename 
    save_path = os.path.join(IMAGES_FOLDER, safe_filename)

    try:
        async with aiofiles.open(save_path, 'wb') as out_file:
            content = await file.read()
            await out_file.write(content)
        print(f"Image '{safe_filename}' uploaded successfully to {IMAGES_FOLDER}")
        return JSONResponse(content={"message": f"Image '{safe_filename}' uploaded successfully."}, status_code=201)
    except Exception as e:
        print(f"Error saving uploaded image '{safe_filename}': {e}")
        raise HTTPException(status_code=500, detail=f"Could not save image: {str(e)}")
    finally:
        await file.close()


@app.get("/list_images/", response_model=ImageListResponse)
async def list_images():
    """Returns a list of image filenames available on the server."""
    try:
        images = [f for f in os.listdir(IMAGES_FOLDER) if os.path.isfile(os.path.join(IMAGES_FOLDER, f))]
        return ImageListResponse(images=images)
    except Exception as e:
        print(f"Error listing images in {IMAGES_FOLDER}: {e}")
        raise HTTPException(status_code=500, detail="Could not list server images.")



@app.get("/images/{filename}")
async def get_image(filename: str):
    """Serves an image file from the server's image directory."""
    try:
        
        
        if ".." in filename or filename.startswith(("/", "\\")):
            raise HTTPException(status_code=400, detail="Invalid filename.")

        
        image_dir = Path(IMAGES_FOLDER)
        file_path = (image_dir / filename).resolve()

        
        if not file_path.is_relative_to(image_dir.resolve()):
             raise HTTPException(status_code=403, detail="Access denied.")

        if file_path.is_file():
            
            
            return FileResponse(str(file_path))
        else:
            raise HTTPException(status_code=404, detail="Image not found.")

    except HTTPException as http_exc:
        
        raise http_exc
    except Exception as e:
        print(f"Error serving image '{filename}': {e}")
        raise HTTPException(status_code=500, detail="Internal server error serving image.")




@app.get("/profiles/", response_model=ProfileListResponse)
async def get_profiles():
    """Returns a list of all user profiles."""
    profiles_data = load_profiles()
    profiles_list = [
        Profile(name=name, image_filename=filename) 
        for name, filename in profiles_data.items()
    ]
    return ProfileListResponse(profiles=profiles_list)

@app.post("/profiles/", response_model=ProfileUpdateResponse)
async def add_or_update_profile(name: str = Form(...), file: UploadFile = File(...)):
    """Adds a new profile or updates the image for an existing profile."""
    profiles = load_profiles()
    
    
    allowed_extensions = {".jpg", ".jpeg", ".png", ".gif", ".webp"}
    file_extension = os.path.splitext(file.filename)[1].lower()
    if file_extension not in allowed_extensions:
        raise HTTPException(status_code=400, detail=f"Invalid image file type. Allowed: {', '.join(allowed_extensions)}")

    
    
    safe_base_name = "".join(c for c in name if c.isalnum() or c in (' ', '-')).rstrip()
    safe_base_name = safe_base_name.replace(' ', '_') 
    if not safe_base_name: 
        safe_base_name = f"profile_{uuid.uuid4().hex[:8]}"
    
    new_filename = f"{safe_base_name}{file_extension}"
    save_path = os.path.join(IMAGES_FOLDER, new_filename)

    
    old_filename = profiles.get(name)

    try:
        
        async with aiofiles.open(save_path, 'wb') as out_file:
            content = await file.read()
            await out_file.write(content)
        print(f"Saved new image for profile '{name}' to {new_filename}")

        
        profiles[name] = new_filename
        save_profiles(profiles)
        message = "Profile added successfully."
        if old_filename and old_filename != new_filename: 
            message = "Profile image updated successfully."
            
            old_file_path = os.path.join(IMAGES_FOLDER, old_filename)
            if os.path.exists(old_file_path):
                try:
                    os.remove(old_file_path)
                    print(f"Deleted old image file: {old_filename}")
                except OSError as e:
                    print(f"Warning: Could not delete old image file {old_filename}: {e}")

        return ProfileUpdateResponse(
            message=message,
            profile=Profile(name=name, image_filename=new_filename)
        )

    except Exception as e:
        print(f"Error saving profile image for '{name}': {e}")
        
        if os.path.exists(save_path):
            try: os.remove(save_path) 
            except OSError: pass
        raise HTTPException(status_code=500, detail=f"Could not save profile image: {str(e)}")
    finally:
        await file.close()

@app.delete("/profiles/{name}", response_model=ProfileDeleteResponse)
async def delete_profile(name: str):
    """Deletes a profile and its associated image."""
    profiles = load_profiles()
    
    filename_to_delete = profiles.get(name)
    
    if filename_to_delete is None:
        raise HTTPException(status_code=404, detail=f"Profile '{name}' not found.")

    
    del profiles[name]
    save_profiles(profiles)

    
    file_path = os.path.join(IMAGES_FOLDER, filename_to_delete)
    if os.path.exists(file_path):
        try:
            os.remove(file_path)
            print(f"Deleted profile image file: {filename_to_delete}")
        except OSError as e:
            print(f"Warning: Could not delete profile image file {filename_to_delete}: {e}")
            

    return ProfileDeleteResponse(message="Profile deleted successfully.", name=name)



if __name__ == "__main__":
    import uvicorn
    print(f"Starting server...")
    print(f"Serving images from: {IMAGES_FOLDER}")
    print(f"Allowed frontend origins: {origins}")
    
    
    
    uvicorn.run("backend_main:app", host="0.0.0.0", port=8000, reload=True)